import hashlib
import os
import re
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, cast
from sqlalchemy import select, desc
from sqlalchemy.ext.asyncio import AsyncSession
from docx import Document as DocxDocument
from openpyxl import Workbook  # type: ignore
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import letter  # type: ignore
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle  # type: ignore
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # type: ignore
from reportlab.lib import colors  # type: ignore

from backend.app.models.user import User
from backend.app.models.deliverable import Deliverable
from backend.app.models.audit import AuditLog
from backend.app.core.document_security import (
    sanitize_filename as core_sanitize_filename,
    assert_path_confined
)

class DeliverableService:
    """
    Sovereign Enterprise Deliverable Generation Engine.
    Produces native DOCX, XLSX, PPTX, TXT, and PDF files stored in on-premises
    isolated scratch storage with integrity verification.
    """

    BASE_DELIVERABLES_DIR = os.environ.get(
        "DELIVERABLES_DIR",
        os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../../scratch/deliverables"))
    )

    def __init__(self):
        os.makedirs(self.BASE_DELIVERABLES_DIR, exist_ok=True)

    def sanitize_filename(self, filename: str, ext: str) -> str:
        """Sanitizes filename against path traversal and dangerous characters using core security."""
        cleaned = core_sanitize_filename(filename)
        if not cleaned.lower().endswith(f".{ext.lower()}"):
            cleaned = f"{cleaned}.{ext.lower()}"
        return cleaned

    def compute_sha256(self, file_path: str) -> str:
        """Calculates cryptographic SHA-256 digest of a local deliverable."""
        hasher = hashlib.sha256()
        with open(file_path, "rb") as f:
            while chunk := f.read(65536):
                hasher.update(chunk)
        return hasher.hexdigest()

    async def generate_docx(
        self,
        title: str,
        sections: List[Dict[str, str]],
        filename: Optional[str] = None,
        author: str = "Sovereign AI Enclave",
        user: Optional[User] = None,
        run_id: Optional[str] = None,
        db: Optional[AsyncSession] = None
    ) -> Deliverable:
        fname = self.sanitize_filename(filename or title.lower().replace(" ", "_"), "docx")
        dest_path = os.path.join(self.BASE_DELIVERABLES_DIR, fname)

        doc = DocxDocument()
        doc.add_heading(title, level=0)
        p_sub = doc.add_paragraph(f"Generated on-premises: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Author: {author}")
        p_sub.runs[0].font.size = Pt(9)
        p_sub.runs[0].font.italic = True

        for sec in sections:
            doc.add_heading(sec.get("heading", "Section"), level=1)
            doc.add_paragraph(sec.get("body", ""))

        doc.save(dest_path)
        return await self._register_deliverable(
            fname=fname,
            fpath=dest_path,
            ftype="DOCX",
            title=title,
            desc=f"Word deliverable containing {len(sections)} sections.",
            user=user,
            run_id=run_id,
            meta={"sections_count": len(sections), "author": author},
            db=db
        )

    async def generate_xlsx(
        self,
        title: str,
        sheet_name: str,
        headers: List[str],
        rows: List[List[Any]],
        filename: Optional[str] = None,
        user: Optional[User] = None,
        run_id: Optional[str] = None,
        db: Optional[AsyncSession] = None
    ) -> Deliverable:
        fname = self.sanitize_filename(filename or title.lower().replace(" ", "_"), "xlsx")
        dest_path = os.path.join(self.BASE_DELIVERABLES_DIR, fname)

        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet(title=sheet_name[:30])
        else:
            ws.title = sheet_name[:30]
        ws.append(headers)
        for r in rows:
            ws.append(r)
        wb.save(dest_path)

        return await self._register_deliverable(
            fname=fname,
            fpath=dest_path,
            ftype="XLSX",
            title=title,
            desc=f"Excel workbook with {len(rows)} data records.",
            user=user,
            run_id=run_id,
            meta={"headers": headers, "rows_count": len(rows)},
            db=db
        )

    async def generate_pptx(
        self,
        title: str,
        slides: List[Dict[str, Any]],
        filename: Optional[str] = None,
        user: Optional[User] = None,
        run_id: Optional[str] = None,
        db: Optional[AsyncSession] = None
    ) -> Deliverable:
        fname = self.sanitize_filename(filename or title.lower().replace(" ", "_"), "pptx")
        dest_path = os.path.join(self.BASE_DELIVERABLES_DIR, fname)

        prs = Presentation()
        # Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        t_shape = cast(Any, slide.shapes.title)
        if t_shape is not None:
            t_shape.text = title
        if len(slide.placeholders) > 1:
            ph1 = cast(Any, slide.placeholders[1])
            if hasattr(ph1, "text_frame"):
                ph1.text_frame.text = f"Sovereign Executive Briefing\n{datetime.now(timezone.utc).strftime('%B %d, %Y')}"
            else:
                ph1.text = f"Sovereign Executive Briefing\n{datetime.now(timezone.utc).strftime('%B %d, %Y')}"

        # Content Slides
        bullet_layout = prs.slide_layouts[1]
        for s in slides:
            c_slide = prs.slides.add_slide(bullet_layout)
            ct_shape = cast(Any, c_slide.shapes.title)
            if ct_shape is not None:
                ct_shape.text = s.get("title", "Slide")
            ph = cast(Any, c_slide.placeholders[1])
            tf = ph.text_frame
            points = s.get("bullet_points", [])
            for idx, pt in enumerate(points):
                if idx == 0:
                    tf.text = pt
                else:
                    p = tf.add_paragraph()
                    p.text = pt

        prs.save(dest_path)

        return await self._register_deliverable(
            fname=fname,
            fpath=dest_path,
            ftype="PPTX",
            title=title,
            desc=f"PowerPoint presentation deck with {len(slides) + 1} slides.",
            user=user,
            run_id=run_id,
            meta={"slides_count": len(slides) + 1},
            db=db
        )

    async def generate_txt(
        self,
        title: str,
        content: str,
        filename: Optional[str] = None,
        user: Optional[User] = None,
        run_id: Optional[str] = None,
        db: Optional[AsyncSession] = None
    ) -> Deliverable:
        fname = self.sanitize_filename(filename or title.lower().replace(" ", "_"), "txt")
        dest_path = os.path.join(self.BASE_DELIVERABLES_DIR, fname)

        with open(dest_path, "w", encoding="utf-8") as f:
            f.write(f"=== {title} ===\n")
            f.write(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}\n\n")
            f.write(content)

        return await self._register_deliverable(
            fname=fname,
            fpath=dest_path,
            ftype="TXT",
            title=title,
            desc="Plain text sovereign deliverable.",
            user=user,
            run_id=run_id,
            meta={"char_count": len(content)},
            db=db
        )

    async def generate_pdf(
        self,
        title: str,
        paragraphs: List[str],
        filename: Optional[str] = None,
        user: Optional[User] = None,
        run_id: Optional[str] = None,
        db: Optional[AsyncSession] = None
    ) -> Deliverable:
        fname = self.sanitize_filename(filename or title.lower().replace(" ", "_"), "pdf")
        dest_path = os.path.join(self.BASE_DELIVERABLES_DIR, fname)

        doc = SimpleDocTemplate(dest_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#0f172a'),
            spaceAfter=14
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155'),
            spaceAfter=10
        )

        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 10))

        for p in paragraphs:
            story.append(Paragraph(p, body_style))

        doc.build(story)

        return await self._register_deliverable(
            fname=fname,
            fpath=dest_path,
            ftype="PDF",
            title=title,
            desc=f"Sovereign PDF document with {len(paragraphs)} sections.",
            user=user,
            run_id=run_id,
            meta={"paragraphs_count": len(paragraphs)},
            db=db
        )

    async def _register_deliverable(
        self,
        fname: str,
        fpath: str,
        ftype: str,
        title: str,
        desc: str,
        user: Optional[User],
        run_id: Optional[str],
        meta: Dict[str, Any],
        db: Optional[AsyncSession]
    ) -> Deliverable:
        file_size = os.path.getsize(fpath)
        sha256 = self.compute_sha256(fpath)

        deliverable = Deliverable(
            filename=fname,
            file_path=fpath,
            file_type=ftype,
            file_size_bytes=file_size,
            sha256_hash=sha256,
            title=title,
            description=desc,
            owner_id=user.id if user else None,
            company_code=getattr(user, "company_code", None) if user else None,
            run_id=run_id,
            status="READY",
            metadata_json=meta
        )

        if db:
            db.add(deliverable)
            if user:
                db.add(AuditLog(
                    actor_id=user.id,
                    actor_email=getattr(user, "email", None),
                    action="DELIVERABLE_CREATED",
                    resource_type="DELIVERABLE",
                    resource_id=deliverable.id,
                    status="SUCCESS",
                    ip_address="127.0.0.1",
                    details={"filename": fname, "file_type": ftype, "sha256": sha256, "file_size_bytes": file_size}
                ))
            await db.commit()
            await db.refresh(deliverable)

        return deliverable

deliverable_service = DeliverableService()
